"""
Build 18-dubai-world-restructuring/Dubai_World_Restructuring.xlsx — the 2010\u201311 restructuring of Dubai World and
Nakheel rebuilt as a restructuring model: the liquidity problem, the recovery each creditor class received in
present-value terms, and the government's debt-for-equity conversion.

Sheets: Cover · Summary · Inputs · Liquidity · Recoveries · Debt-for-Equity · Sensitivity · Checks

On 25 November 2009 Dubai World asked its creditors for a standstill, and for eighteen months the largest
restructuring the Gulf had seen worked through three questions a restructuring banker always asks. Can the
business meet its maturities from cash and asset sales (liquidity)? If not, what does each class of creditor
actually receive once terms are extended, coupons cut and paper substituted for cash (recovery)? And who takes the
equity risk that the creditors will not (debt-for-equity)? The model rebuilds each from the terms announced in
May 2010 and approved in September 2010, and from Nakheel's separate plan of 2011.

Source data: Dubai World and Government of Dubai announcements of 20 May and 10 September 2010; Nakheel
announcements of 2010\u20132011; contemporaneous reporting of the tranche terms. Figures are as publicly reported;
where a term was reported as a range, the model states the value used.
"""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from ibkit.style import (Book, Theme, F_INPUT, F_FORMULA, F_LINK, F_TEXT, F_BOLD, F_NOTE,
                         NUM, NUM0, NUM2, PCT, MULT, TOTAL_BORDER, DOUBLE_BORDER, recalc, export_pdf, preview_png)
from ibkit.sheet import Sheet, checks_sheet

HERE = os.path.dirname(os.path.abspath(__file__)); FONT = "Arial"
DW = Theme("dubaiworld", "Dubai World", "1F3A5F", "6B7A8F", "C8A951", "EEF1F6", tagline="Restructuring 2010\u201311")

book = Book(theme=DW, project_no=18, project="Restructuring: Liquidity, Recoveries and Debt-for-Equity",
            company="Dubai World and Nakheel", units="US$ billions unless stated", as_of="terms of May\u2013Sep 2010 and Nakheel 2011")
S = lambda name, **kw: Sheet(book, name, **kw)

# --------------------------------------------------------------------------- #
# Inputs
# --------------------------------------------------------------------------- #
si = S("Inputs", ncols=6, label_width=66, subtitle="The announced terms, and the stated assumptions needed to value them")
si.ws.column_dimensions["F"].width = 62
I = {}
si.section("DUBAI WORLD \u2014 CLAIMS AT THE STANDSTILL (US$ bn)")
I["bank"] = si.row("Bank and financial creditors (about 70 lenders)", 14.4, NUM, bold=True)
I["gov_claims"] = si.row("Government of Dubai claims (existing loans and the Dubai Financial Support Fund advances)", 8.9, NUM)
I["other"] = si.row("Other claims included in the restructuring perimeter", 1.6, NUM, note="Total perimeter reported as about US$24.9bn")
I["total"] = si.row("Total restructured", f"={I['bank']}+{I['gov_claims']}+{I['other']}", NUM, bold=True, border=TOTAL_BORDER)
si.blank(); si.section("BANK TERMS AGREED (September 2010)")
I["A_amt"] = si.row("Tranche A: amount", 4.4, NUM, bold=True)
I["A_ten"] = si.row("Tranche A: tenor (years, bullet)", 5, '0')
I["A_cpn"] = si.row("Tranche A: cash coupon", 0.01, PCT)
I["B_amt"] = si.row("Tranche B: amount", 10.0, NUM, bold=True)
I["B_ten"] = si.row("Tranche B: tenor (years, bullet)", 8, '0')
I["B_cpn"] = si.row("Tranche B: cash coupon", 0.01, PCT)
I["B_pik"] = si.row("Tranche B: payment-in-kind margin (option chosen by most US$ lenders)", 0.015, PCT,
                    note="Lenders chose between 1.5% PIK with a government shortfall guarantee, 2.5% with partial cover, and 3.5% with none")
I["B_guar"] = si.row("Share of Tranche B covered by the government shortfall guarantee (stated)", 0.60, PCT)
I["orig_cpn"] = si.row("Original contractual coupon on the bank debt (stated: pre-crisis LIBOR plus margin, all-in)", 0.055, PCT)
si.blank(); si.section("GOVERNMENT OF DUBAI")
I["conv"] = si.row("Claims converted to equity in Dubai World", f"={I['gov_claims']}", NUM, font=F_LINK, bold=True)
I["new_money"] = si.row("New money committed: working capital facility", 0.5, NUM)
I["int_support"] = si.row("New money committed: interest support over the term", 1.0, NUM)
si.blank(); si.section("NAKHEEL (separated from Dubai World; its own plan, 2011)")
I["nk_trade"] = si.row("Trade creditor claims", 2.7, NUM, note="About AED 10bn of contractor and supplier claims")
I["nk_trade_cash"] = si.row("  paid in cash", 0.40, PCT)
I["nk_trade_sukuk"] = si.row("  paid in a 5-year sukuk", 0.60, PCT)
I["nk_sukuk_cpn"] = si.row("  sukuk profit rate", 0.10, PCT, note="AED 4.8bn sukuk issued August 2011 at 10%")
I["nk_bank"] = si.row("Nakheel bank debt extended", 1.6, NUM)
I["nk_bank_ten"] = si.row("  extension (years)", 5, '0')
I["nk_bank_cpn"] = si.row("  coupon after extension (stated)", 0.04, PCT)
I["nk_gov"] = si.row("Government support injected into Nakheel (equity and converted debt)", 8.0, NUM, bold=True)
si.blank(); si.section("VALUATION ASSUMPTIONS \u2014 stated")
I["disc"] = si.row("Discount rate for a Dubai government-related credit in 2010 (yield on comparable paper)", 0.08, PCT,
                   note="Dubai sovereign five-year paper yielded 6\u20137% in 2010; corporate GREs traded wider")
I["guar_spread"] = si.row("Spread by which the guaranteed portion is discounted below the corporate rate", 0.015, PCT)
I["disc_guar"] = si.row("Discount rate for the government-guaranteed portion", f"={I['disc']}-{I['guar_spread']}", PCT, font=F_FORMULA)
I["sukuk_disc"] = si.row("Discount rate for the Nakheel sukuk (traded near par to a small discount after issue)", 0.11, PCT)
si.blank(); si.section("ASSET REALISATION PLAN (US$ bn, as described in the plan)")
I["dpw"] = si.row("Stake in DP World (about 80%, at 2010 market value)", 8.5, NUM)
I["istithmar"] = si.row("Istithmar World portfolio (Barneys, hotels, financial stakes)", 3.0, NUM)
I["drydocks"] = si.row("Drydocks World and Dubai Maritime City", 1.5, NUM)
I["jafza"] = si.row("Economic Zones World (Jafza)", 4.0, NUM)
I["other_assets"] = si.row("Other assets and cash generation over the plan", 2.0, NUM)
I["haircut"] = si.row("Realisation haircut on non-DP World assets in a forced sale", 0.30, PCT)
I["dpw_haircut"] = si.row("Realisation haircut on the DP World stake (listed; block discount)", 0.15, PCT)

# --------------------------------------------------------------------------- #
# Liquidity
# --------------------------------------------------------------------------- #
sl = S("Liquidity", ncols=6, label_width=66, subtitle="Why the standstill happened, and whether the eight-year plan could be met from assets")
sl.ws.column_dimensions["F"].width = 58
Lq = {}
sl.section("THE PROBLEM IN NOVEMBER 2009")
Lq["nk_sukuk_09"] = sl.row("Nakheel sukuk maturing 14 December 2009 (US$ bn)", 3.52, NUM, bold=True)
Lq["cash_09"] = sl.row("Group cash available to meet it without support (stated)", 0.5, NUM)
Lq["gap_09"] = sl.row("Shortfall on a single maturity", f"={Lq['nk_sukuk_09']}-{Lq['cash_09']}", NUM, bold=True, border=TOTAL_BORDER)
Lq["abu_dhabi"] = sl.row("Abu Dhabi support that paid it (14 December 2009, US$ bn)", 10.0, NUM, note="Of which US$4.1bn to Nakheel's sukuk and the rest to the standstill period")
Lq["mat_2010_11"] = sl.row("Further group maturities falling in 2010\u20132011 before the plan (stated)", 6.0, NUM)
Lq["cover"] = sl.row("Cash plus Abu Dhabi support divided by maturities to end-2011", f"=({Lq['cash_09']}+{Lq['abu_dhabi']})/({Lq['nk_sukuk_09']}+{Lq['mat_2010_11']})", MULT, bold=True)
sl.note("A liquidity crisis is a mismatch of dates, not a shortage of assets. Dubai World's assets were worth more than its debt on "
        "most estimates; it simply could not turn a port operator and a property portfolio into US$3.5bn in three weeks. The "
        "standstill bought time; Abu Dhabi's US$10bn bought the first maturity.", height=44)
sl.blank(); sl.section("CAN THE EIGHT-YEAR PLAN BE MET FROM ASSET SALES?")
Lq["gross"] = sl.row("Gross asset value in the plan", f"={I['dpw']}+{I['istithmar']}+{I['drydocks']}+{I['jafza']}+{I['other_assets']}", NUM0, bold=True)
Lq["net"] = sl.row("Net realisable after haircuts", f"={I['dpw']}*(1-{I['dpw_haircut']})+({I['istithmar']}+{I['drydocks']}+{I['jafza']}+{I['other_assets']})*(1-{I['haircut']})", NUM, bold=True)
Lq["A_due"] = sl.row("Tranche A due in year five", f"={I['A_amt']}", NUM, font=F_LINK)
Lq["B_due"] = sl.row("Tranche B due in year eight, including accrued PIK", f"={I['B_amt']}*(1+{I['B_pik']})^{I['B_ten']}", NUM)
Lq["total_due"] = sl.row("Total bank principal to repay over the plan", f"={Lq['A_due']}+{Lq['B_due']}", NUM, bold=True, border=TOTAL_BORDER)
Lq["asset_cover"] = sl.row("Net realisable assets / bank principal", f"={Lq['net']}/{Lq['total_due']}", MULT, bold=True, border=DOUBLE_BORDER)
Lq["cash_int"] = sl.row("Annual cash interest under the new terms (both tranches at 1%)", f"=({I['A_amt']}+{I['B_amt']})*0.01", NUM2)
Lq["cash_int_old"] = sl.row("  compared with cash interest at the original all-in coupon", f"=({I['A_amt']}+{I['B_amt']})*{I['orig_cpn']}", NUM2)
Lq["int_saving"] = sl.row("Annual cash interest saved by the coupon cut", f"={Lq['cash_int_old']}-{Lq['cash_int']}", NUM2, bold=True)
Lq["no_dpw"] = sl.row("Asset cover without selling DP World (the crown jewel the plan hoped to keep)", f"=({Lq['net']}-{I['dpw']}*(1-{I['dpw_haircut']}))/{Lq['total_due']}", MULT)
sl.note("On forced-sale 2010 values the plan did not quite cover the debt even with DP World, and covered less than half without it. "
        "That is the case for the eight-year tenor: the plan relied on asset values recovering, and on DP World growing into the debt, "
        "rather than on sales into the trough. The 1% coupon \u2014 US$0.65bn a year of cash interest saved \u2014 is what made waiting affordable.", height=48)

# --------------------------------------------------------------------------- #
# Recoveries
# --------------------------------------------------------------------------- #
sr = S("Recoveries", ncols=6, label_width=66, subtitle="What each class received, in present value \u2014 the nominal number was 100 cents; the economic one was not")
sr.ws.column_dimensions["F"].width = 58
R = {}
def pv_bullet(amt, cpn, ten, disc, pik="0"):
    # PV of a bullet paying cash coupon annually and PIK accreting to principal
    return f"={amt}*({cpn}*(1-(1+{disc})^-{ten})/{disc}+(1+{pik})^{ten}/(1+{disc})^{ten})"
sr.section("DUBAI WORLD BANK CREDITORS")
R["A_pv"] = sr.row("Tranche A: PV at the discount rate (US$ bn)", pv_bullet(I["A_amt"], I["A_cpn"], I["A_ten"], I["disc"]), NUM2, bold=True)
R["A_rec"] = sr.row("  recovery, cents on the dollar", f"={R['A_pv']}/{I['A_amt']}", PCT, bold=True)
R["B_g_pv"] = sr.row("Tranche B, guaranteed portion: PV at the guaranteed rate", pv_bullet(f"{I['B_amt']}*{I['B_guar']}", I["B_cpn"], I["B_ten"], I["disc_guar"], I["B_pik"]), NUM2)
R["B_u_pv"] = sr.row("Tranche B, unguaranteed portion: PV at the corporate rate", pv_bullet(f"{I['B_amt']}*(1-{I['B_guar']})", I["B_cpn"], I["B_ten"], I["disc"], I["B_pik"]), NUM2)
R["B_pv"] = sr.row("Tranche B: PV", f"={R['B_g_pv']}+{R['B_u_pv']}", NUM2, bold=True)
R["B_rec"] = sr.row("  recovery, cents on the dollar", f"={R['B_pv']}/{I['B_amt']}", PCT, bold=True)
R["bank_pv"] = sr.row("Bank creditors: total PV", f"={R['A_pv']}+{R['B_pv']}", NUM2, bold=True, border=TOTAL_BORDER)
R["bank_rec"] = sr.row("Bank creditors: blended recovery", f"={R['bank_pv']}/{I['bank']}", PCT, bold=True, border=DOUBLE_BORDER)
R["bank_haircut"] = sr.row("  economic haircut (1 \u2212 recovery)", f"=1-{R['bank_rec']}", PCT, bold=True)
R["old_pv"] = sr.row("For comparison: PV of the original debt at its original coupon over five years at the same discount rate", pv_bullet(I["bank"], I["orig_cpn"], 5, I["disc"]), NUM2)
R["old_rec"] = sr.row("  as cents on the dollar (what the loans were worth even before the restructuring, at 2010 yields)", f"={R['old_pv']}/{I['bank']}", PCT)
R["transfer"] = sr.row("Value transferred from banks to the company by the new terms (US$ bn)", f"={R['old_pv']}-{R['bank_pv']}", NUM2, bold=True)
sr.blank(); sr.section("NAKHEEL CREDITORS")
R["nk_cash"] = sr.row("Trade creditors: cash element", f"={I['nk_trade']}*{I['nk_trade_cash']}", NUM2)
R["nk_sukuk_pv"] = sr.row("Trade creditors: PV of the sukuk element", pv_bullet(f"{I['nk_trade']}*{I['nk_trade_sukuk']}", I["nk_sukuk_cpn"], 5, I["sukuk_disc"]), NUM2)
R["nk_trade_rec"] = sr.row("Trade creditors: recovery", f"=({R['nk_cash']}+{R['nk_sukuk_pv']})/{I['nk_trade']}", PCT, bold=True, border=DOUBLE_BORDER)
R["nk_bank_pv"] = sr.row("Nakheel banks: PV of the extended loan", pv_bullet(I["nk_bank"], I["nk_bank_cpn"], I["nk_bank_ten"], I["disc"]), NUM2)
R["nk_bank_rec"] = sr.row("Nakheel banks: recovery", f"={R['nk_bank_pv']}/{I['nk_bank']}", PCT, bold=True)
sr.blank(); sr.section("THE WATERFALL, IN ONE TABLE")
sr.head(["Claim", "PV received", "Recovery"], label="Class (ranked by recovery)")
rows = [("Nakheel trade creditors (40% cash, 60% 10% sukuk)", I["nk_trade"], f"{R['nk_cash']}+{R['nk_sukuk_pv']}"),
        ("Nakheel banks (5-year extension)", I["nk_bank"], R["nk_bank_pv"]),
        ("Dubai World Tranche A (5 years, 1%)", I["A_amt"], R["A_pv"]),
        ("Dubai World Tranche B (8 years, 1% + PIK, part-guaranteed)", I["B_amt"], R["B_pv"]),
        ("Government of Dubai (converted to equity)", I["gov_claims"], "0")]
r0 = sr.r
for name, claim, pv in rows:
    ws = sr.ws; r = sr.r
    ws.cell(r, 2, name).font = F_TEXT
    for col, v, f in [(3, f"={claim}", NUM2), (4, f"={pv}", NUM2)]:
        c = ws.cell(r, col, v); c.font = F_LINK if col == 3 else F_FORMULA; c.number_format = f
    c = ws.cell(r, 5, f"=D{r}/C{r}"); c.font = F_FORMULA; c.number_format = PCT
    sr.r += 1
R["wf_r0"], R["wf_r1"] = r0, sr.r - 1
sr.note("Nobody took a nominal haircut except the government, which took all of it by converting to equity. The banks' 'full recovery' "
        "was worth roughly 70 cents on the dollar in present value: five to eight years at 1% when the market required 8%. Trade "
        "creditors, who could stop building, did best. That ordering \u2014 operational leverage over legal seniority \u2014 is the lesson of "
        "most Gulf restructurings of the period.", height=52)

# --------------------------------------------------------------------------- #
# Debt-for-Equity
# --------------------------------------------------------------------------- #
se = S("Debt-for-Equity", ncols=6, label_width=66, subtitle="The government's conversion: what it gave up, what it received, and why the banks would not do it")
se.ws.column_dimensions["F"].width = 58
E = {}
se.section("THE CONVERSION")
E["conv"] = se.row("Government claims converted to Dubai World equity (US$ bn)", f"={I['conv']}", NUM, font=F_LINK, bold=True)
E["new"] = se.row("New money committed alongside", f"={I['new_money']}+{I['int_support']}", NUM, font=F_LINK)
E["total_gov"] = se.row("Total government exposure to the plan", f"={E['conv']}+{E['new']}", NUM, bold=True, border=TOTAL_BORDER)
se.blank(); se.section("WHAT THE EQUITY WAS WORTH")
E["ev"] = se.row("Enterprise value of the restructured group: net realisable assets (Liquidity sheet)", f"={Lq['net']}", NUM, font=F_LINK)
E["debt_pv"] = se.row("Less: bank debt at present value (the market value of the claims ranking ahead)", f"=-{R['bank_pv']}", NUM, font=F_LINK)
E["debt_face"] = se.row("  memo: bank debt at face", f"=-{I['bank']}", NUM, font=F_LINK)
E["eq_pv"] = se.row("Implied equity value on a present-value basis", f"={E['ev']}+{E['debt_pv']}", NUM, bold=True, border=DOUBLE_BORDER)
E["eq_face"] = se.row("Implied equity value if the debt is taken at face", f"={E['ev']}+{E['debt_face']}", NUM, bold=True)
E["gov_rec_pv"] = se.row("Government recovery on the converted claims: equity value / claims converted (PV basis)", f"=MAX(0,{E['eq_pv']})/{E['conv']}", PCT, bold=True)
E["gov_rec_face"] = se.row("  on a face-value basis", f"=MAX(0,{E['eq_face']})/{E['conv']}", PCT)
E["breakeven"] = se.row("Asset value at which the government's equity is worth its converted claims (US$ bn)", f"={R['bank_pv']}+{E['conv']}", NUM, bold=True,
                        note="With DP World's later performance the group cleared this level; in 2010 it was a bet")
se.blank(); se.section("WHY THE BANKS TOOK PAPER AND THE GOVERNMENT TOOK EQUITY")
E["bank_alt"] = se.row("If the banks had converted 30% of their claims to equity instead: equity claim they would hold (US$ bn)", f"={I['bank']}*0.3", NUM)
E["bank_alt_pv"] = se.row("  PV of the remaining 70% on the new terms plus a pro-rata share of the PV equity", f"={R['bank_pv']}*0.7+{E['bank_alt']}/({E['bank_alt']}+{E['conv']})*MAX(0,{E['ev']}-{R['bank_pv']}*0.7)", NUM2)
E["bank_alt_rec"] = se.row("  recovery under that alternative", f"={E['bank_alt_pv']}/{I['bank']}", PCT, bold=True)
E["vs"] = se.row("  versus the recovery they accepted", f"={E['bank_alt_rec']}-{R['bank_rec']}", PCT, bold=True)
se.note("The government converted because it was the only creditor whose objective was not recovery: it needed the group to survive "
        "and its name to be honoured. The banks took paper because paper with a 1% coupon and a sovereign shortfall guarantee could "
        "stay on the books at close to par, while equity in a Dubai holding company in 2010 could not. Regulatory accounting, not "
        "valuation, chose the instruments.", height=52)

# --------------------------------------------------------------------------- #
# Sensitivity
# --------------------------------------------------------------------------- #
ss = S("Sensitivity", ncols=8, label_width=40, col_width=13, subtitle="Bank recovery under different discount rates and Tranche B guarantee cover")
DISCS = [0.06, 0.07, 0.08, 0.09, 0.10]; GUARS = [0.0, 0.3, 0.6, 0.8, 1.0]
ss.section("BLENDED BANK RECOVERY: DISCOUNT RATE (rows) versus GUARANTEED SHARE OF TRANCHE B (columns)")
ss.ws.cell(ss.r, 2, "\u2193 discount rate  \\  guaranteed share \u2192").font = F_BOLD
for j, g in enumerate(GUARS):
    c = ss.ws.cell(ss.r, 3 + j, g); c.font = book.f_header; c.fill = book.fill_secondary; c.number_format = PCT; c.alignment = Alignment(horizontal="center")
h = ss.r; ss.r += 1; g0 = ss.r
def rec_formula(disc, guar):
    a = f"{I['A_amt']}*({I['A_cpn']}*(1-(1+{disc})^-{I['A_ten']})/{disc}+1/(1+{disc})^{I['A_ten']})"
    dg = f"({disc}-{I['guar_spread']})"
    bg = f"{I['B_amt']}*{guar}*({I['B_cpn']}*(1-(1+{dg})^-{I['B_ten']})/{dg}+(1+{I['B_pik']})^{I['B_ten']}/(1+{dg})^{I['B_ten']})"
    bu = f"{I['B_amt']}*(1-{guar})*({I['B_cpn']}*(1-(1+{disc})^-{I['B_ten']})/{disc}+(1+{I['B_pik']})^{I['B_ten']}/(1+{disc})^{I['B_ten']})"
    return f"=({a}+{bg}+{bu})/{I['bank']}"
for d in DISCS:
    c = ss.ws.cell(ss.r, 2, d); c.font = F_BOLD; c.number_format = PCT
    for j, g in enumerate(GUARS):
        cc = ss.ws.cell(ss.r, 3 + j, rec_formula(f"$B{ss.r}", f"{L(3+j)}${h}")); cc.font = F_FORMULA; cc.number_format = PCT
        if abs(d - 0.08) < 1e-9 and abs(g - 0.6) < 1e-9: cc.fill = book.fill_accent; cc.font = Font(name=FONT, size=10, bold=True)
    ss.r += 1
ss.blank()
ss.note("At 2010 yields the banks' recovery sat between 65 and 75 cents whatever the guarantee cover; only a discount rate near the "
        "1% coupon itself \u2014 which is what a bank's accounting effectively assumed \u2014 takes it back to par. The grid is the whole argument "
        "between the economists and the accountants.", height=40)

# --------------------------------------------------------------------------- #
# Summary
# --------------------------------------------------------------------------- #
sy = S("Summary", ncols=6, label_width=66, subtitle="The restructuring in numbers")
sy.ws.column_dimensions["F"].width = 56
sy.section("LIQUIDITY")
sy.row("Shortfall on the December 2009 Nakheel sukuk (US$ bn)", f"={Lq['gap_09']}", NUM, font=F_LINK, bold=True)
sy.row("Net realisable assets / bank principal over the plan", f"={Lq['asset_cover']}", MULT, font=F_LINK, bold=True)
sy.row("  without DP World", f"={Lq['no_dpw']}", MULT, font=F_LINK)
sy.row("Annual cash interest saved by the coupon cut (US$ bn)", f"={Lq['int_saving']}", NUM2, font=F_LINK)
sy.blank(); sy.section("RECOVERIES, PRESENT VALUE")
sy.row("Tranche A (5 years, 1%)", f"={R['A_rec']}", PCT, font=F_LINK)
sy.row("Tranche B (8 years, 1% + PIK, 60% guaranteed)", f"={R['B_rec']}", PCT, font=F_LINK)
sy.row("Bank creditors, blended", f"={R['bank_rec']}", PCT, font=F_LINK, bold=True)
sy.row("Value transferred from banks to the company (US$ bn)", f"={R['transfer']}", NUM2, font=F_LINK, bold=True)
sy.row("Nakheel trade creditors", f"={R['nk_trade_rec']}", PCT, font=F_LINK)
sy.row("Nakheel banks", f"={R['nk_bank_rec']}", PCT, font=F_LINK)
sy.blank(); sy.section("DEBT-FOR-EQUITY")
sy.row("Government claims converted (US$ bn)", f"={E['conv']}", NUM, font=F_LINK, bold=True)
sy.row("Implied equity value, PV basis (US$ bn)", f"={E['eq_pv']}", NUM, font=F_LINK)
sy.row("Government recovery on conversion, PV basis", f"={E['gov_rec_pv']}", PCT, font=F_LINK, bold=True)
sy.row("Asset value at which the conversion breaks even (US$ bn)", f"={E['breakeven']}", NUM, font=F_LINK, border=DOUBLE_BORDER)
sy.blank()
sy.bullets([
    "The crisis was a maturity, not a balance sheet: US$3.5bn due in three weeks against US$0.5bn of cash, in a group whose assets "
    "covered its debt. Every term that followed \u2014 five and eight-year bullets, 1% coupons, an asset-sale programme \u2014 was designed to "
    "buy time for DP World to grow into the debt rather than be sold into the trough. On 2010 forced-sale values the assets covered "
    "about 93% of the bank principal; the plan was a bet on recovery, and the tenor was the size of the bet.",
    "Nobody took a nominal haircut except the government. The banks' '100 cents' was worth about 70 in present value at 2010 yields; "
    "the coupon cut alone transferred several billion dollars from lenders to the company. Regulatory accounting let the banks carry "
    "the paper near par, which is why they accepted it and why they would not take equity.",
    "Trade creditors did best. Nakheel's contractors received 40% in cash and 60% in a 10% sukuk that traded near par: a recovery "
    "above the banks' despite ranking below them, because they could stop the projects. Operational leverage beat legal seniority.",
    "The government's conversion was the only true risk capital in the plan. On 2010 asset values its equity was worth a fraction of "
    "the US$8.9bn converted; it broke even only if the assets reached the bank debt's present value plus its own claim. DP World's "
    "subsequent performance is why the bet paid, not the terms.",
])

# --------------------------------------------------------------------------- #
# Checks
# --------------------------------------------------------------------------- #
tests = [
    ("Restructured perimeter reconciles to the reported US$24.9bn within US$0.1bn", f"=ABS({I['total']}-24.9)<0.1"),
    ("Tranche A plus Tranche B equal the bank claims", f"=ABS({I['A_amt']}+{I['B_amt']}-{I['bank']})<0.001"),
    ("Tranche A PV reproduces a hand calculation (1% five-year bullet at 8%)", f"=ABS({R['A_pv']}/{I['A_amt']}-(0.01*(1-1.08^-5)/0.08+1.08^-5))<0.0001"),
    ("Tranche B recovery is below Tranche A recovery (longer, riskier)", f"={R['B_rec']}<{R['A_rec']}"),
    ("Blended bank recovery lies between 55 and 85 cents", f"=AND({R['bank_rec']}>0.55,{R['bank_rec']}<0.85)"),
    ("Value transfer equals PV of old terms less PV of new terms", f"=ABS({R['transfer']}-({R['old_pv']}-{R['bank_pv']}))<0.0001"),
    ("Nakheel trade recovery exceeds the banks' (operational leverage)", f"={R['nk_trade_rec']}>{R['bank_rec']}"),
    ("Nakheel cash and sukuk shares sum to 100%", f"=ABS({I['nk_trade_cash']}+{I['nk_trade_sukuk']}-1)<0.0001"),
    ("Waterfall table recoveries are all between 0% and 100%", f"=AND(MIN(Recoveries!E{R['wf_r0']}:E{R['wf_r1']})>=0,MAX(Recoveries!E{R['wf_r0']}:E{R['wf_r1']})<=1)"),
    ("Asset cover on forced-sale 2010 values is close to one times with DP World (0.85\u20131.15x) and well below without it (<0.7x)", f"=AND({Lq['asset_cover']}>0.85,{Lq['asset_cover']}<1.15,{Lq['no_dpw']}<0.7)"),
    ("Government break-even exceeds net realisable assets on 2010 values (it was a bet)", f"={E['breakeven']}>{Lq['net']}"),
    ("Base sensitivity cell reconciles to the Recoveries sheet", f"=ABS(Sensitivity!E{g0+2}-{R['bank_rec']})<0.0001"),
    ("Recovery falls as the discount rate rises in every column of the grid", "=AND(" + ",".join(f"Sensitivity!{L(3+j)}{g0+4}<Sensitivity!{L(3+j)}{g0}" for j in range(5)) + ")"),
    ("Recovery rises with guarantee cover in every row of the grid", "=AND(" + ",".join(f"Sensitivity!G{g0+i}>=Sensitivity!C{g0+i}" for i in range(5)) + ")"),
]
checks_sheet(book, tests,
             "The third check reproduces a tranche present value by hand. The rest test the relationships the case is built to "
             "show: recovery falls with tenor and rises with guarantee cover, trade creditors beat banks, and the plan needed DP World.")

# --------------------------------------------------------------------------- #
book.cover(
    blurb="The 2010\u201311 restructuring of Dubai World and Nakheel rebuilt as a restructuring model from the announced terms: the "
          "liquidity mismatch that forced the November 2009 standstill, whether the eight-year plan could be met from asset sales, "
          "the present-value recovery of each creditor class under the extended low-coupon terms, and the Government of Dubai's "
          "US$8.9bn debt-for-equity conversion.",
    method=[
        "Liquidity: the December 2009 maturity against available cash; net realisable assets after stated haircuts against the "
        "two bank tranches at maturity, with and without the DP World stake; cash interest saved by the coupon cut.",
        "Recoveries: each tranche valued as a bullet with its cash coupon, PIK accretion and guarantee cover at a stated 2010 "
        "discount rate; Nakheel's 40% cash / 60% sukuk trade settlement and extended bank debt on the same basis; a ranked waterfall.",
        "Debt-for-equity: the equity value implied by assets less the present value of bank claims, the government's recovery, "
        "its break-even asset value, and the counterfactual in which the banks had taken equity.",
        "A sensitivity of bank recovery to the discount rate and to guarantee cover.",
    ],
    toc=[("Summary", "liquidity, recoveries, debt-for-equity"),
         ("Liquidity", "the 2009 maturity; asset cover over the plan"),
         ("Recoveries", "PV by tranche and class; the waterfall"),
         ("Debt-for-Equity", "the government's conversion and break-even"),
         ("Sensitivity", "discount rate against guarantee cover"),
         ("Inputs", "announced terms and stated assumptions"),
         ("Checks", "fourteen tests; must read MODEL OK")],
    highlights=[("Shortfall on the Dec 2009 sukuk (US$ bn)", f"={Lq['gap_09']}", NUM),
                ("Asset cover with / without DP World", f"=TEXT({Lq['asset_cover']},\"0.00\")&\"x / \"&TEXT({Lq['no_dpw']},\"0.00\")&\"x\"", "@"),
                ("Bank recovery, present value", f"={R['bank_rec']}", PCT),
                ("Value transferred from banks (US$ bn)", f"={R['transfer']}", NUM2),
                ("Nakheel trade creditor recovery", f"={R['nk_trade_rec']}", PCT),
                ("Government recovery on conversion, PV", f"={E['gov_rec_pv']}", PCT)],
    sources=["Dubai World and Government of Dubai announcements of 20 May 2010 (terms proposed) and 10 September 2010 (creditor approval); Nakheel announcements of 2010\u20132011 including the August 2011 sukuk; Abu Dhabi support of 14 December 2009.",
             "Tranche amounts, tenors, coupons and PIK options, the government conversion and Nakheel's cash/sukuk split are as publicly reported. Discount rates, asset values, haircuts and the guarantee share are stated inputs.",
             "The model is a reconstruction for learning from public information; it is not a statement about any creditor's actual position or accounting."])

book.finish()
path = os.path.join(HERE, "Dubai_World_Restructuring.xlsx")
book.save(path)
print("saved", path)
print("recalc:", recalc(path))
pdf = export_pdf(path)
print("pdf:", pdf, os.path.exists(pdf))
png = preview_png(pdf, page=0, dpi=80)
os.replace(png, os.path.join(HERE, "cover.png"))
print("cover.png written")
