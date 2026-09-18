"""
ibkit.sheet — a running-row-cursor helper on top of ibkit.style.Book.

Projects 6 and 7 each carried an identical copy of this class inside build_model.py. From Project 8 onward
it lives here. The idea is unchanged: every value written through `row` or `multi` records its absolute
cell reference in `self.ref`, so later formulas are assembled from names rather than hand-counted cells,
and a row inserted anywhere cannot silently break a formula elsewhere.

    s = Sheet(book, "Inputs", ncols=6)
    s.section("FINANCIALS")
    rev = s.row("Revenue", 1000.0, NUM0)            # returns "'Inputs'!$C$5"
    ebitda = s.row("EBITDA", "=" + rev + "*0.2")     # formula font applied automatically
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L

from .style import F_INPUT, F_FORMULA, F_TEXT, F_BOLD, F_NOTE, NUM, FONT


class Sheet:
    def __init__(self, book, name, ncols, label_width=52, col_width=15, subtitle=None, valcol=3):
        self.book = book
        self.ws = book.sheet(name, ncols=ncols, label_width=label_width, col_width=col_width, subtitle=subtitle)
        self.name, self.ncols, self.valcol, self.r = name, ncols, valcol, 4
        self.ref = {}

    # -- structure -----------------------------------------------------------
    def section(self, text):
        self.book.section(self.ws, self.r, text, self.ncols); self.r += 1

    def subsection(self, text):
        self.book.subsection(self.ws, self.r, text, self.ncols); self.r += 1

    def blank(self, n=1):
        self.r += n

    def head(self, labels, first_col=None, label=""):
        self.book.year_header(self.ws, self.r, labels, first_col=first_col or self.valcol, label=label); self.r += 1

    def note(self, text, height=30):
        self.book.note(self.ws, self.r, text, self.ncols, height=height); self.r += 1

    def bullets(self, texts):
        for t in texts:
            self.note("\u2022  " + t, height=15 * max(2, len(t) // 100 + 1))

    # -- values --------------------------------------------------------------
    def row(self, label, value, fmt=NUM, font=None, note="", bold=False, border=None, col=None, key=None):
        ws, r = self.ws, self.r
        c = ws.cell(r, 2, label); c.font = F_BOLD if bold else F_TEXT
        col = col or self.valcol
        v = ws.cell(r, col, value)
        if font is None:
            font = F_FORMULA if (isinstance(value, str) and value.startswith("=")) else F_INPUT
        v.font = font; v.number_format = fmt
        if border: v.border = border
        if note:
            n = ws.cell(r, self.ncols, note); n.font = F_NOTE; n.alignment = Alignment(wrap_text=True, vertical="top")
        ref = f"'{self.name}'!${L(col)}${r}"
        self.ref[key or label] = ref
        self.r += 1
        return ref

    def multi(self, label, values, fmt=NUM, font=None, bold=False, border=None, first_col=None, rel=False, key=None, note=""):
        """One label, several columns. Returns a list of refs (absolute unless rel=True)."""
        ws, r = self.ws, self.r
        c = ws.cell(r, 2, label); c.font = F_BOLD if bold else F_TEXT
        fc = first_col or self.valcol
        if note:
            n = ws.cell(r, self.ncols, note); n.font = F_NOTE; n.alignment = Alignment(wrap_text=True, vertical="top")
        refs = []
        for i, value in enumerate(values):
            v = ws.cell(r, fc + i, value)
            f = font or (F_FORMULA if (isinstance(value, str) and value.startswith("=")) else F_INPUT)
            v.font = f; v.number_format = fmt
            if border: v.border = border
            refs.append(f"{L(fc + i)}{r}" if rel else f"'{self.name}'!${L(fc + i)}${r}")
        self.ref[key or label] = refs
        self.r += 1
        return refs

    def text_row(self, label, texts, first_col=None, bold=False, wrap=False):
        """A row of plain text cells (for tables with narrative columns)."""
        ws, r = self.ws, self.r
        ws.cell(r, 2, label).font = F_BOLD if bold else F_TEXT
        fc = first_col or self.valcol
        for i, t in enumerate(texts):
            c = ws.cell(r, fc + i, t); c.font = F_BOLD if bold else F_TEXT
            if wrap: c.alignment = Alignment(wrap_text=True, vertical="top")
        self.r += 1

    def status_cell(self, label, formula):
        """The MODEL OK / CHECK FAILED cell used on every Checks sheet."""
        ws = self.ws
        ws.cell(self.r, 2, label).font = Font(name=FONT, size=12, bold=True, color=self.book.theme.primary)
        v = ws.cell(self.r, 3, formula)
        v.font = Font(name=FONT, size=12, bold=True); v.fill = self.book.fill_accent
        v.alignment = Alignment(horizontal="center")
        self.r += 1


def checks_sheet(book, tests, closing_note="", name="Checks"):
    """Build the standard Checks sheet from a list of (label, boolean formula) pairs."""
    sk = Sheet(book, name, ncols=5, label_width=74, col_width=16, subtitle="Every check must read TRUE")
    sk.section("INTEGRITY CHECKS")
    for label, fml in tests:
        sk.ws.cell(sk.r, 2, label).font = F_TEXT
        v = sk.ws.cell(sk.r, 3, fml); v.font = F_FORMULA; v.alignment = Alignment(horizontal="center")
        sk.r += 1
    k0, k1 = 5, sk.r - 1
    sk.blank()
    sk.status_cell("MODEL STATUS", f'=IF(COUNTIF(C{k0}:C{k1},FALSE)=0,"MODEL OK","CHECK FAILED")')
    if closing_note:
        sk.blank(); sk.note(closing_note, height=34)
    return sk
